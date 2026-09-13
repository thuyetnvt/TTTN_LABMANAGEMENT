from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"N:\thuyet-dev\LabManagementtt\docs\HUONG_DAN_CAI_DAT_VA_TRIEN_KHAI.docx"
REPOSITORY_URL = "https://github.com/thuyetnvt/TTTN_LABMANAGEMENT.git"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="000000", size=10, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[index], "1F4E78")
        if widths:
            table.rows[0].cells[index].width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[index], value, align=align)
            if widths:
                cells[index].width = Cm(widths[index])
            if row_index % 2:
                set_cell_shading(cells[index], "F3F6F9")
    document.add_paragraph()
    return table


def add_code(document, code):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.right_indent = Cm(0.45)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    for index, line in enumerate(code.splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def add_step(document, number, title, body):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"Bước {number}. {title}. ")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    body_run = paragraph.add_run(body)
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(11)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

title_style = styles["Title"]
title_style.font.name = "Times New Roman"
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("HƯỚNG DẪN CÀI ĐẶT VÀ TRIỂN KHAI LABMANAGEMENT")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Dành cho người tiếp nhận mã nguồn và quản trị hệ thống").italic = True
doc.add_paragraph()

intro = doc.add_paragraph()
intro.add_run("Mục đích. ").bold = True
intro.add_run(
    "Tài liệu hướng dẫn tải mã nguồn từ GitHub, chạy hệ thống bằng Docker, kiểm tra sau cài đặt, "
    "cập nhật phiên bản và triển khai lên VPS. Cách chạy bằng Docker Compose được ưu tiên vì không cần "
    "cài riêng .NET, Node.js và MySQL."
)

doc.add_heading("1. Thông tin dự án", level=1)
add_table(doc, ["Nội dung", "Thông tin"], [
    ["Tên hệ thống", "LabManagement - Quản lý tài sản và thiết bị Phòng Lab IoT"],
    ["Mã nguồn", REPOSITORY_URL],
    ["Frontend", "Vue 3, Vite, Ant Design Vue"],
    ["Backend", "ASP.NET Core Web API, Entity Framework Core"],
    ["Cơ sở dữ liệu", "MySQL 8.4"],
    ["Cách chạy khuyến nghị", "Docker Compose v2"],
], widths=[4.3, 12.2])

doc.add_heading("2. Chuẩn bị", level=1)
for item in (
    "Máy tính đã cài Git và Docker Desktop; Docker Desktop phải ở trạng thái Running.",
    "Có quyền đọc repository GitHub. Nếu repository riêng tư, đăng nhập GitHub trước khi tải.",
    "Bảo đảm các cổng dự kiến sử dụng chưa bị ứng dụng khác chiếm dụng.",
    "Không sao chép mật khẩu, JWT key hoặc file .env lên GitHub."
):
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("3. Tải mã nguồn từ GitHub", level=1)
doc.add_heading("3.1. Tải bằng Git", level=2)
add_step(doc, 1, "Mở PowerShell", "Chọn thư mục sẽ lưu dự án.")
add_step(doc, 2, "Clone repository", "Chạy các lệnh sau:")
add_code(doc, f"git clone {REPOSITORY_URL}\ncd TTTN_LABMANAGEMENT")
add_step(doc, 3, "Kiểm tra nhánh", "Nhánh dùng để chạy chính thức là main.")
add_code(doc, "git branch --show-current\ngit status")

doc.add_heading("3.2. Tải bằng Download ZIP", level=2)
doc.add_paragraph(
    "Mở trang GitHub của dự án, chọn Code, chọn Download ZIP, giải nén rồi mở PowerShell tại thư mục vừa giải nén. "
    "Cách này phù hợp để xem hoặc chạy thử; nên dùng Git nếu cần cập nhật phiên bản thường xuyên."
)

doc.add_heading("4. Chạy hệ thống trên máy cá nhân", level=1)
add_step(doc, 1, "Tạo file cấu hình", "Sao chép file mẫu thành .env.")
add_code(doc, "Copy-Item .env.example .env")
add_step(doc, 2, "Đổi các giá trị bảo mật", "Mở .env và thay tối thiểu các biến sau:")
add_code(doc, "MYSQL_ROOT_PASSWORD=<mat_khau_root_manh>\nMYSQL_PASSWORD=<mat_khau_ung_dung_manh>\nJWT_KEY=<chuoi_ngau_nhien_toi_thieu_32_ky_tu>\nJWT_REFRESH_TOKEN_DAYS=7\nSEED_ENABLED=true\nSEED_DEFAULT_PASSWORD=<mat_khau_demo_manh>")
doc.add_paragraph(
    "SEED_ENABLED=true chỉ dùng khi cần tạo tài khoản mẫu lần đầu. Sau khi đăng nhập và đổi mật khẩu, đặt lại thành false."
)
add_step(doc, 3, "Kiểm tra cấu hình Docker", "Lệnh phải kết thúc mà không báo lỗi.")
add_code(doc, "docker compose config --quiet")
add_step(doc, 4, "Build và khởi động", "Lần đầu có thể mất vài phút để tải image và build ứng dụng.")
add_code(doc, "docker compose up -d --build\ndocker compose ps")
add_step(doc, 5, "Mở hệ thống", "Truy cập địa chỉ mặc định http://localhost:8080.")

doc.add_heading("5. Kiểm tra sau khi cài đặt", level=1)
add_table(doc, ["STT", "Nội dung kiểm tra", "Kết quả mong đợi"], [
    ["1", "docker compose ps", "Các service ở trạng thái Up hoặc healthy"],
    ["2", "Mở http://localhost:8080", "Trang đăng nhập hiển thị bình thường"],
    ["3", "Đăng nhập tài khoản mẫu", "Hệ thống yêu cầu đổi mật khẩu lần đầu"],
    ["4", "Mở thiết bị, mượn trả, báo cáo", "Không có lỗi tải dữ liệu hoặc lỗi 403 sai quyền"],
    ["5", "Mở /health qua frontend", "Backend và database phản hồi khỏe"],
], widths=[1.2, 7.3, 8.0])

doc.add_heading("6. Các lệnh vận hành thường dùng", level=1)
add_table(doc, ["Mục đích", "Lệnh"], [
    ["Xem trạng thái", "docker compose ps"],
    ["Xem log backend", "docker compose logs -f backend"],
    ["Xem log frontend", "docker compose logs -f frontend"],
    ["Khởi động lại", "docker compose restart backend frontend"],
    ["Dừng nhưng giữ dữ liệu", "docker compose down"],
    ["Build lại sau khi sửa code", "docker compose up -d --build"],
], widths=[6.2, 10.3])
warning = doc.add_paragraph()
warning.add_run("Lưu ý quan trọng. ").bold = True
warning.add_run("Không chạy docker compose down -v trên hệ thống có dữ liệu thật vì tùy chọn -v xóa các volume dữ liệu.")

doc.add_heading("7. Cập nhật mã nguồn", level=1)
doc.add_paragraph("Trước khi cập nhật, phải sao lưu database, file upload và khóa Data Protection.")
add_code(doc, "git status\ngit pull --ff-only origin main\ndocker compose up -d --build\ndocker compose ps")
doc.add_paragraph(
    "Nếu git status hiển thị file đã sửa trên máy chủ, không tự xóa hoặc ghi đè. Cần sao lưu thay đổi đó và xử lý cùng người quản trị repository."
)

doc.add_heading("8. Triển khai trên VPS", level=1)
doc.add_heading("8.1. Kiểm tra quyền truy cập", level=2)
add_code(doc, "ssh -p <SSH_PORT> <USERNAME>@<SERVER_IP>\ncd /hdd1/lab\ndocker ps")
doc.add_paragraph(
    "Nếu docker ps hiển thị danh sách container thì tài khoản đã có quyền đọc Docker. Nếu báo permission denied, liên hệ quản trị viên máy chủ để cấp quyền; không tự dùng sudo khi chưa được phép."
)
doc.add_heading("8.2. Triển khai bằng GitHub Actions", level=2)
doc.add_paragraph(
    "Repository đã có workflow CI. Khi push lên nhánh triển khai, hệ thống chạy kiểm thử backend, frontend, E2E và build Docker. "
    "Job Deploy VPS chỉ chạy khi các job trước thành công và GitHub Secrets đã được cấu hình."
)
add_table(doc, ["Secret", "Ý nghĩa"], [
    ["VPS_HOST", "Địa chỉ máy chủ"],
    ["VPS_USER", "Tài khoản SSH"],
    ["VPS_PORT", "Cổng SSH"],
    ["VPS_SSH_KEY", "Private key dùng riêng cho CI"],
    ["VPS_DEPLOY_DIR", "Thư mục dự án trên VPS; máy chủ hiện tại dùng /hdd1/lab"],
    ["SMTP_USERNAME / SMTP_PASSWORD", "Tài khoản gửi email nếu bật SMTP"],
    ["GOOGLE_CLIENT_ID", "Client ID nếu bật đăng nhập Google"],
], widths=[6.2, 10.3])
add_code(doc, "git add <cac_file_da_kiem_tra>\ngit commit -m \"Mo ta noi dung cap nhat\"\ngit push origin main")
doc.add_paragraph(
    "Chỉ xác nhận triển khai thành công khi toàn bộ workflow màu xanh, docker compose ps trên VPS ổn định và trang /health phản hồi bình thường."
)

doc.add_heading("8.3. Kiểm tra bản Git sạch trước khi đẩy", level=2)
doc.add_paragraph(
    "Chạy kiểm tra từ thư mục gốc. Chỉ tạo commit khi toàn bộ lệnh kết thúc thành công và danh sách file không chứa secret, file cấu hình local hoặc kết quả build."
)
add_code(
    doc,
    "git status --short\n"
    "git diff --check\n"
    "dotnet test .\\lab-backend.Tests\\LabManagementAPI.Tests.csproj --configuration Release\n"
    "Set-Location .\\lab-frontend\n"
    "npm ci\n"
    "npm test\n"
    "npm run build\n"
    "Set-Location .."
)
add_code(
    doc,
    "git add -A\n"
    "git diff --cached --check\n"
    "git diff --cached --stat\n"
    "git commit -m \"Mo ta noi dung cap nhat\"\n"
    "git status --short\n"
    "git push origin main"
)
doc.add_paragraph(
    "Sau khi commit, git status --short phải không in ra dòng nào. Nếu còn file ngoài phạm vi bàn giao, dừng lại và kiểm tra trước khi push."
)

doc.add_page_break()
doc.add_heading("9. Xử lý lỗi thường gặp", level=1)
add_table(doc, ["Hiện tượng", "Cách xử lý"], [
    ["SSH connection timed out", "Kiểm tra lại mạng, IP, cổng SSH và firewall phía VPS; dùng Test-NetConnection <IP> -Port <PORT>."],
    ["docker ps báo permission denied", "Tài khoản chưa có quyền Docker; liên hệ quản trị viên máy chủ."],
    ["Cổng 8080 đã được sử dụng", "Đổi APP_PORT trong .env hoặc dừng ứng dụng đang chiếm cổng."],
    ["Backend không kết nối MySQL", "Kiểm tra MYSQL_DATABASE, MYSQL_USER, MYSQL_PASSWORD và log service db/backend."],
    ["Đăng nhập xong bị trả về trang login", "Kiểm tra JWT_KEY, thời gian máy chủ và migration bảng RefreshTokens."],
    ["Migration lỗi", "Sao lưu database, đọc log migration; không sửa schema production thủ công khi chưa có bản sao."],
], widths=[5.7, 10.8])

doc.add_heading("10. Checklist bàn giao", level=1)
for item in (
    "Mã nguồn đã được push lên GitHub và workflow kiểm thử đạt.",
    "File .env và các secret không nằm trong commit.",
    "Database, upload và Data Protection keys đã có bản sao lưu.",
    "Năm vai trò đã được kiểm tra: Admin, Trưởng lab, Phó lab, Giảng viên và Sinh viên.",
    "Luồng mượn, duyệt, bàn giao, xác nhận nhận, trả và cấp phát đã được thử.",
    "Import/export Excel, báo cáo PDF, thông báo và quên mật khẩu đã được kiểm tra.",
    "Domain HTTPS và /health hoạt động trước khi nghiệm thu."
):
    doc.add_paragraph("☐ " + item)

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("LabManagement - Hướng dẫn cài đặt và triển khai")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)

doc.save(OUTPUT)
print(OUTPUT)
